import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib.pyplot as plt
import numpy as np
import math

def exportar_para_excel(bateria_resultados, nome_arquivo="tabelas_comparativas.xlsx"):
    """
    Gera uma planilha Excel (.xlsx) profissional contendo as tabelas comparativas
    de todos os problemas resolvidos, pronta para ser copiada para um documento Word/Docs.
    """
    wb = openpyxl.Workbook()
    # Remove a folha padrão
    wb.remove(wb.active)

    # Cores e estilos
    cor_cabecalho = "2F4F4F"       # Cinza ardósia escuro
    cor_fundo_sub = "E8EEF0"       # Fundo claro
    fonte_branca = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fonte_negrito = Font(name="Calibri", size=11, bold=True)
    fonte_normal = Font(name="Calibri", size=11)
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_esquerda = Alignment(horizontal="left", vertical="center")
    alinhamento_direita = Alignment(horizontal="right", vertical="center")

    borda_fina = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    # 1. Folha Resumo Geral (Todas as tabelas reunidas)
    ws_geral = wb.create_sheet(title="Tabelas Consolidadas")
    ws_geral.views.sheetView[0].showGridLines = True

    linha_atual = 2

    for prob in bateria_resultados:
        nome_prob = prob["nome"]
        dados = prob["dados"]

        # Título do Problema
        ws_geral.merge_cells(start_row=linha_atual, start_column=2, end_row=linha_atual, end_column=7)
        celula_titulo = ws_geral.cell(row=linha_atual, column=2, value=nome_prob.upper())
        celula_titulo.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        celula_titulo.fill = PatternFill(start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid")
        celula_titulo.alignment = alinhamento_centro
        ws_geral.row_dimensions[linha_atual].height = 28
        linha_atual += 1

        # Cabeçalhos da Tabela
        colunas = ["Método", "Raiz Encontrada", "Iterações", "Tempo (ms)", "Precisão |f(raiz)|", "Status"]
        for col_idx, col_nome in enumerate(colunas, start=2):
            cell = ws_geral.cell(row=linha_atual, column=col_idx, value=col_nome)
            cell.font = fonte_negrito
            cell.fill = PatternFill(start_color=cor_fundo_sub, end_color=cor_fundo_sub, fill_type="solid")
            cell.alignment = alinhamento_centro
            cell.border = borda_fina
        ws_geral.row_dimensions[linha_atual].height = 22
        linha_atual += 1

        # Dados das Linhas
        for item in dados:
            c_metodo = ws_geral.cell(row=linha_atual, column=2, value=item["metodo"])
            c_metodo.alignment = alinhamento_esquerda
            c_metodo.font = fonte_normal
            c_metodo.border = borda_fina

            c_raiz = ws_geral.cell(row=linha_atual, column=3, value=round(item["raiz"], 8) if item["raiz"] is not None else "Falha")
            c_raiz.alignment = alinhamento_direita
            c_raiz.font = fonte_normal
            c_raiz.border = borda_fina
            if isinstance(c_raiz.value, float):
                c_raiz.number_format = "0.00000000"

            c_iter = ws_geral.cell(row=linha_atual, column=4, value=item["iteracoes"])
            c_iter.alignment = alinhamento_centro
            c_iter.font = fonte_normal
            c_iter.border = borda_fina

            c_tempo = ws_geral.cell(row=linha_atual, column=5, value=round(item["tempo_ms"], 4))
            c_tempo.alignment = alinhamento_direita
            c_tempo.font = fonte_normal
            c_tempo.border = borda_fina
            c_tempo.number_format = "0.0000"

            c_res = ws_geral.cell(row=linha_atual, column=6, value=item["residuo"] if item["residuo"] is not None else "-")
            c_res.alignment = alinhamento_direita
            c_res.font = fonte_normal
            c_res.border = borda_fina
            if isinstance(c_res.value, float):
                c_res.number_format = "0.00E+00"

            c_status = ws_geral.cell(row=linha_atual, column=7, value=item["status"])
            c_status.alignment = alinhamento_centro
            c_status.font = fonte_normal
            c_status.border = borda_fina

            ws_geral.row_dimensions[linha_atual].height = 20
            linha_atual += 1

        linha_atual += 2  # Espaçamento entre tabelas

    # Ajusta largura das colunas
    larguras = {2: 18, 3: 18, 4: 12, 5: 14, 6: 18, 7: 18}
    for col_idx, larg in larguras.items():
        col_letter = get_column_letter(col_idx)
        ws_geral.column_dimensions[col_letter].width = larg

    wb.save(nome_arquivo)
    print(f"Planilha Excel salva com sucesso em: {nome_arquivo}")


def gerar_graficos(bateria_resultados, pasta_saida="graficos"):
    """
    Gera gráficos comparativos em alta resolução (.png) salvos em uma pasta,
    prontos para serem inseridos no relatório final em formato Word/Docs.
    """
    os.makedirs(pasta_saida, exist_ok=True)
    plt.style.use('seaborn-v0_8-whitegrid' if 'seaborn-v0_8-whitegrid' in plt.style.available else 'default')

    # Cores padronizadas para os métodos
    cores_metodos = {
        "Bissecção": "#3498db",      # Azul
        "Falsa Posição": "#e67e22",  # Laranja
        "Newton-Raphson": "#2ecc71", # Verde
        "Secante": "#9b59b6"         # Roxo
    }

    # 1. Gráfico de Comparação de Iterações
    nomes_problemas = [p["nome"].split(":")[0] for p in bateria_resultados]
    metodos = ["Bissecção", "Falsa Posição", "Newton-Raphson", "Secante"]

    x = np.arange(len(nomes_problemas))
    largura = 0.2

    fig, ax = plt.subplots(figsize=(9, 5), dpi=300)

    max_iter_val = 0
    for i, metodo in enumerate(metodos):
        iteracoes = []
        for prob in bateria_resultados:
            val = next((m["iteracoes"] for m in prob["dados"] if m["metodo"] == metodo), 0)
            iteracoes.append(val)
            if val > max_iter_val:
                max_iter_val = val
        
        posicoes = x + (i - 1.5) * largura
        barras = ax.bar(posicoes, iteracoes, largura, label=metodo, color=cores_metodos[metodo])

        # Adiciona rótulo de valor em cima de cada barra
        for barra in barras:
            altura = barra.get_height()
            ax.annotate(f'{altura}',
                        xy=(barra.get_x() + barra.get_width() / 2, altura),
                        xytext=(0, 3), textcoords="offset points",
                        ha='center', va='bottom', fontsize=9, fontweight='bold')

    ax.set_ylabel('Número de Iterações', fontsize=12, fontweight='bold')
    ax.set_title('Comparação do Número de Iterações por Método Numérico', fontsize=14, fontweight='bold', pad=15)
    ax.set_xticks(x)
    ax.set_xticklabels(nomes_problemas, fontsize=11)
    ax.set_ylim(0, max_iter_val * 1.35)
    ax.legend(title='Método', loc='upper right', frameon=True, facecolor='white', edgecolor='#cccccc', framealpha=1.0, fontsize=10, title_fontsize=10)
    ax.grid(axis='y', linestyle='--', alpha=0.7)
    plt.tight_layout()

    caminho_iter = os.path.join(pasta_saida, "comparativo_iteracoes.png")
    plt.savefig(caminho_iter, dpi=300)
    plt.close()
    print(f"Gráfico de iterações salvo em: {caminho_iter}")

    # 2. Gráfico de Comparação de Tempo de Execução
    fig, ax = plt.subplots(figsize=(9, 5), dpi=300)

    max_tempo_val = 0.0
    for i, metodo in enumerate(metodos):
        tempos = []
        for prob in bateria_resultados:
            val = next((m["tempo_ms"] for m in prob["dados"] if m["metodo"] == metodo), 0.0)
            tempos.append(val)
            if val > max_tempo_val:
                max_tempo_val = val
        
        posicoes = x + (i - 1.5) * largura
        barras = ax.bar(posicoes, tempos, largura, label=metodo, color=cores_metodos[metodo])

        for barra in barras:
            altura = barra.get_height()
            ax.annotate(f'{altura:.3f}',
                        xy=(barra.get_x() + barra.get_width() / 2, altura),
                        xytext=(0, 3), textcoords="offset points",
                        ha='center', va='bottom', fontsize=8)

    ax.set_ylabel('Tempo de Execução (ms)', fontsize=12, fontweight='bold')
    ax.set_title('Comparação do Tempo de Execução (ms)', fontsize=14, fontweight='bold', pad=15)
    ax.set_xticks(x)
    ax.set_xticklabels(nomes_problemas, fontsize=11)
    ax.set_ylim(0, max_tempo_val * 1.35)
    ax.legend(title='Método', loc='upper right', frameon=True, facecolor='white', edgecolor='#cccccc', framealpha=1.0, fontsize=10, title_fontsize=10)
    ax.grid(axis='y', linestyle='--', alpha=0.7)
    plt.tight_layout()

    caminho_tempo = os.path.join(pasta_saida, "comparativo_tempo.png")
    plt.savefig(caminho_tempo, dpi=300)
    plt.close()
    print(f"Gráfico de tempos salvo em: {caminho_tempo}")

    # 3. Gráficos das Funções Matemáticas e suas Raízes (Bonificação do PDF)
    fig, axs = plt.subplots(1, 3, figsize=(15, 4.5), dpi=300)

    # 3.1 Exemplo: f(x) = cos(x) - x
    x1 = np.linspace(0, 1.2, 300)
    y1 = np.cos(x1) - x1
    raiz1 = 0.73908513
    axs[0].plot(x1, y1, color="#2980b9", lw=2, label="f(x) = cos(x) - x")
    axs[0].axhline(0, color="gray", linestyle="--", lw=1)
    axs[0].plot(raiz1, 0, "ro", markersize=8, label=f"Raiz ≈ {raiz1:.4f}")
    axs[0].set_title("Exemplo de Teste", fontsize=11, fontweight="bold")
    axs[0].set_xlabel("x")
    axs[0].set_ylabel("f(x)")
    axs[0].set_ylim(-1.0, 1.5)
    axs[0].legend(loc='upper right', frameon=True, facecolor='white', edgecolor='#cccccc', framealpha=1.0, fontsize=9)
    axs[0].grid(True, linestyle="--", alpha=0.6)

    # 3.2 Problema 1: Resfriamento
    x2 = np.linspace(0, 15, 300)
    y2 = 20.0 + 90.0 * np.exp(-0.15 * x2) - 50.0
    raiz2 = 7.32408192
    axs[1].plot(x2, y2, color="#c0392b", lw=2, label="f(t) = 20+90e^(-0.15t)-50")
    axs[1].axhline(0, color="gray", linestyle="--", lw=1)
    axs[1].plot(raiz2, 0, "ro", markersize=8, label=f"t ≈ {raiz2:.4f} min")
    axs[1].set_title("Problema 1: Resfriamento", fontsize=11, fontweight="bold")
    axs[1].set_xlabel("t (minutos)")
    axs[1].set_ylabel("f(t)")
    axs[1].set_ylim(-40, 85)
    axs[1].legend(loc='upper right', frameon=True, facecolor='white', edgecolor='#cccccc', framealpha=1.0, fontsize=9)
    axs[1].grid(True, linestyle="--", alpha=0.6)

    # 3.3 Problema 2: Capacitor RC
    x3 = np.linspace(0, 4, 300)
    y3 = 5.0 * (1.0 - np.exp(-x3)) - 3.8
    raiz3 = 1.42711636
    axs[2].plot(x3, y3, color="#27ae60", lw=2, label="f(t) = 5(1-e^(-t))-3.8")
    axs[2].axhline(0, color="gray", linestyle="--", lw=1)
    axs[2].plot(raiz3, 0, "ro", markersize=8, label=f"t ≈ {raiz3:.4f} s")
    axs[2].set_title("Problema 2: Circuito RC", fontsize=11, fontweight="bold")
    axs[2].set_xlabel("t (segundos)")
    axs[2].set_ylabel("f(t)")
    axs[2].set_ylim(-4.2, 2.2)
    axs[2].legend(loc='upper right', frameon=True, facecolor='white', edgecolor='#cccccc', framealpha=1.0, fontsize=9)
    axs[2].grid(True, linestyle="--", alpha=0.6)

    plt.suptitle("Funções Analisadas e Localização das Raízes (Zeros)", fontsize=14, fontweight="bold", y=1.02)
    plt.tight_layout()

    caminho_curvas = os.path.join(pasta_saida, "curvas_funcoes_e_zeros.png")
    plt.savefig(caminho_curvas, dpi=300, bbox_inches='tight')
    plt.close()
    print(f"Gráfico das funções e zeros salvo em: {caminho_curvas}")
